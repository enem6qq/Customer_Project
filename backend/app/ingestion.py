"""Dokumenten-Ingestion.

Liest Dokumente aus dem Ablageverzeichnis ein und zerlegt sie in Chunks.
Der oberste Ordnername unterhalb des Ablageverzeichnisses ist die
Zugriffsgruppe des Dokuments (z. B. data/documents/personal/urlaub.md
-> Gruppe "personal"). Dateien direkt im Wurzelverzeichnis gehören zur
Gruppe "allgemein".

Unterstützte Formate: .txt, .md, .pdf, .docx, .pptx, .xlsx, .csv, .html
Weitere Quellen (SharePoint, Netzlaufwerke, E-Mail-Postfächer) werden später
als zusätzliche Loader ergänzt – die Chunk-Struktur bleibt gleich.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

logger = logging.getLogger(__name__)

STANDARD_GRUPPE = "allgemein"


@dataclass
class Chunk:
    text: str
    dokument: str          # Anzeige-/Suchname (relativ zum Ablageverzeichnis)
    titel: str             # Dateiname ohne Endung
    gruppe: str            # Zugriffsgruppe (RBAC)
    position: int          # laufende Nummer innerhalb des Dokuments
    pfad_absolut: str = "" # Originaldatei auf der Platte (für Öffnen/Kopieren)
    metadata: dict = field(default_factory=dict)


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_pptx(path: Path) -> str:
    from pptx import Presentation

    praesentation = Presentation(str(path))
    teile = []
    for nummer, folie in enumerate(praesentation.slides, start=1):
        texte = []
        for form in folie.shapes:
            if getattr(form, "has_text_frame", False) and form.text_frame.text.strip():
                texte.append(form.text_frame.text)
            if getattr(form, "has_table", False):
                for zeile in form.table.rows:
                    texte.append(" | ".join(zelle.text for zelle in zeile.cells))
        if folie.has_notes_slide and folie.notes_slide.notes_text_frame.text.strip():
            texte.append("Notizen: " + folie.notes_slide.notes_text_frame.text)
        if texte:
            teile.append(f"[Folie {nummer}]\n" + "\n".join(texte))
    return "\n\n".join(teile)


def _read_xlsx(path: Path) -> str:
    import openpyxl

    arbeitsmappe = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        teile = []
        for blatt in arbeitsmappe.worksheets:
            zeilen = []
            for zeile in blatt.iter_rows(values_only=True):
                zellen = [str(wert) for wert in zeile if wert is not None and str(wert).strip()]
                if zellen:
                    zeilen.append(" | ".join(zellen))
            if zeilen:
                teile.append(f"[Tabellenblatt: {blatt.title}]\n" + "\n".join(zeilen))
        return "\n\n".join(teile)
    finally:
        arbeitsmappe.close()


def _read_csv(path: Path) -> str:
    import csv

    inhalt = path.read_text(encoding="utf-8", errors="replace")
    try:
        dialekt = csv.Sniffer().sniff(inhalt[:2048], delimiters=",;\t")
    except csv.Error:
        dialekt = csv.excel
    zeilen = []
    for zeile in csv.reader(inhalt.splitlines(), dialekt):
        zellen = [z.strip() for z in zeile if z.strip()]
        if zellen:
            zeilen.append(" | ".join(zellen))
    return "\n".join(zeilen)


class _HTMLTextExtractor(HTMLParser):
    """Zieht sichtbaren Text aus HTML (ohne Skripte/Styles) – nur Standardbibliothek."""

    IGNORIEREN = {"script", "style", "noscript", "template"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}

    def __init__(self) -> None:
        super().__init__()
        self._teile: list[str] = []
        self._ignoriere = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.IGNORIEREN:
            self._ignoriere += 1
        elif tag in self.BLOCK:
            self._teile.append("\n")

    def handle_endtag(self, tag):
        if tag in self.IGNORIEREN and self._ignoriere > 0:
            self._ignoriere -= 1
        elif tag in self.BLOCK:
            self._teile.append("\n")

    def handle_data(self, data):
        if self._ignoriere == 0 and data.strip():
            self._teile.append(data)

    @property
    def text(self) -> str:
        roh = "".join(self._teile)
        zeilen = [z.strip() for z in roh.splitlines()]
        return "\n".join(z for z in zeilen if z)


def _read_html(path: Path) -> str:
    parser = _HTMLTextExtractor()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    return parser.text


# Dateien oberhalb dieser Größe werden übersprungen (Schutz vor GB-Videos
# mit Dokumenten-Endung o. Ä.).
MAX_DATEI_MB = 100

LOADERS = {
    ".txt": _read_txt,
    ".md": _read_txt,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".pptx": _read_pptx,
    ".xlsx": _read_xlsx,
    ".csv": _read_csv,
    ".html": _read_html,
    ".htm": _read_html,
}


def split_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    """Zerlegt Text an Absatzgrenzen in Chunks von ca. `chunk_size` Zeichen."""
    absaetze = [a.strip() for a in text.split("\n\n") if a.strip()]
    chunks: list[str] = []
    aktuell = ""
    for absatz in absaetze:
        if aktuell and len(aktuell) + len(absatz) + 2 > chunk_size:
            chunks.append(aktuell)
            # Überlappung: Ende des letzten Chunks als Kontext mitnehmen
            aktuell = aktuell[-overlap:] if overlap > 0 else ""
        aktuell = (aktuell + "\n\n" + absatz).strip() if aktuell else absatz
        # Sehr lange Absätze hart teilen
        while len(aktuell) > chunk_size:
            chunks.append(aktuell[:chunk_size])
            aktuell = aktuell[chunk_size - overlap:] if overlap > 0 else aktuell[chunk_size:]
    if aktuell:
        chunks.append(aktuell)
    return chunks


def lade_dokumente(
    basis: Path | list[Path], chunk_size: int = 800, overlap: int = 150
) -> list[Chunk]:
    basen = basis if isinstance(basis, list) else [basis]
    chunks: list[Chunk] = []
    mehrere_basen = len(basen) > 1

    for basis_pfad in basen:
        if not basis_pfad.exists():
            logger.warning("Dokumentenverzeichnis %s existiert nicht", basis_pfad)
            continue
        for datei in sorted(basis_pfad.rglob("*")):
            if not datei.is_file() or datei.suffix.lower() not in LOADERS:
                continue
            if datei.stat().st_size > MAX_DATEI_MB * 1024 * 1024:
                logger.warning("%s ist größer als %d MB – wird übersprungen", datei, MAX_DATEI_MB)
                continue
            relativ = datei.relative_to(basis_pfad)
            gruppe = relativ.parts[0] if len(relativ.parts) > 1 else STANDARD_GRUPPE
            # Bei mehreren Ablageorten den Ordnernamen voranstellen, damit die
            # Anzeige eindeutig bleibt (z. B. "Vertraege/kunde_a/vertrag.pdf").
            name = f"{basis_pfad.name}/{relativ.as_posix()}" if mehrere_basen else relativ.as_posix()
            try:
                text = LOADERS[datei.suffix.lower()](datei)
            except Exception:
                logger.exception("Konnte %s nicht lesen – wird übersprungen", datei)
                continue
            for i, teil in enumerate(split_text(text, chunk_size, overlap)):
                chunks.append(
                    Chunk(
                        text=teil,
                        dokument=name,
                        titel=datei.stem,
                        gruppe=gruppe,
                        position=i,
                        pfad_absolut=str(datei.resolve()),
                    )
                )
        logger.info("Dokumente aus %s geladen", basis_pfad)
    logger.info("%d Chunks insgesamt", len(chunks))
    return chunks
