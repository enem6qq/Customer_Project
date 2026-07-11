import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.ingestion import lade_dokumente, split_text


def test_split_text_respektiert_chunkgroesse():
    text = "\n\n".join(f"Absatz {i} " + "x" * 120 for i in range(10))
    chunks = split_text(text, chunk_size=300, overlap=50)
    assert len(chunks) > 1
    assert all(len(c) <= 300 for c in chunks)


def test_split_text_teilt_lange_absaetze():
    chunks = split_text("y" * 2000, chunk_size=500, overlap=100)
    assert all(len(c) <= 500 for c in chunks)
    assert sum(len(c) for c in chunks) >= 2000  # nichts geht verloren


def test_gruppe_aus_ordnername(tmp_path):
    (tmp_path / "personal").mkdir()
    (tmp_path / "personal" / "a.md").write_text("Inhalt Personal", encoding="utf-8")
    (tmp_path / "b.md").write_text("Inhalt Wurzel", encoding="utf-8")
    chunks = lade_dokumente(tmp_path)
    gruppen = {c.dokument: c.gruppe for c in chunks}
    assert gruppen["personal/a.md"] == "personal"
    assert gruppen["b.md"] == "allgemein"
    # Absoluter Pfad zur Originaldatei wird mitgeführt (Quellen-Klick)
    assert all(Path(c.pfad_absolut).is_file() for c in chunks)


def test_pptx_loader(tmp_path):
    from pptx import Presentation

    praesentation = Presentation()
    folie = praesentation.slides.add_slide(praesentation.slide_layouts[1])
    folie.shapes.title.text = "SmartWeedControl Pitch"
    folie.placeholders[1].text = "Unkrauterkennung per App"
    datei = tmp_path / "pitch.pptx"
    praesentation.save(str(datei))

    chunks = lade_dokumente(tmp_path)
    text = " ".join(c.text for c in chunks)
    assert "SmartWeedControl Pitch" in text
    assert "Unkrauterkennung" in text
    assert "[Folie 1]" in text


def test_xlsx_loader(tmp_path):
    import openpyxl

    arbeitsmappe = openpyxl.Workbook()
    blatt = arbeitsmappe.active
    blatt.title = "Budget"
    blatt.append(["Posten", "Betrag"])
    blatt.append(["Marketing", 5000])
    arbeitsmappe.save(str(tmp_path / "budget.xlsx"))

    chunks = lade_dokumente(tmp_path)
    text = " ".join(c.text for c in chunks)
    assert "Tabellenblatt: Budget" in text
    assert "Marketing | 5000" in text


def test_csv_loader(tmp_path):
    (tmp_path / "kontakte.csv").write_text(
        "Name;Abteilung\nMüller;Personal\n", encoding="utf-8"
    )
    chunks = lade_dokumente(tmp_path)
    assert "Müller | Personal" in chunks[0].text


def test_html_loader(tmp_path):
    (tmp_path / "seite.html").write_text(
        "<html><head><style>body{color:red}</style><script>alert(1)</script></head>"
        "<body><h1>Anleitung</h1><p>Schritt eins ausführen.</p></body></html>",
        encoding="utf-8",
    )
    chunks = lade_dokumente(tmp_path)
    text = chunks[0].text
    assert "Anleitung" in text
    assert "Schritt eins ausführen." in text
    assert "alert" not in text  # Skripte/Styles werden entfernt
    assert "color:red" not in text


def test_zu_grosse_dateien_werden_uebersprungen(tmp_path, monkeypatch):
    import app.ingestion as ingestion

    monkeypatch.setattr(ingestion, "MAX_DATEI_MB", 0)  # alles ist "zu groß"
    (tmp_path / "riesig.md").write_text("Inhalt", encoding="utf-8")
    assert lade_dokumente(tmp_path) == []


def test_mehrere_ablageorte(tmp_path):
    """Mehrere Ordner werden gemeinsam indiziert; Namen bleiben eindeutig."""
    ablage_a = tmp_path / "Projekte"
    ablage_b = tmp_path / "Vertraege"
    ablage_a.mkdir()
    ablage_b.mkdir()
    (ablage_a / "notiz.md").write_text("Projektnotiz", encoding="utf-8")
    (ablage_b / "notiz.md").write_text("Vertragsnotiz", encoding="utf-8")
    chunks = lade_dokumente([ablage_a, ablage_b])
    namen = sorted(c.dokument for c in chunks)
    assert namen == ["Projekte/notiz.md", "Vertraege/notiz.md"]
